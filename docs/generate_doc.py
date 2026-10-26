"""Generate the StudyFlow AI submission Word document (Deliverable 1).

Run:  backend/.venv/bin/python docs/generate_doc.py
Output: docs/StudyFlow_AI_Submission.docx
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

BRAND = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1E, 0x29, 0x3B)
GREY = RGBColor(0x47, 0x55, 0x69)
LIGHT_BG = "EFF4FF"
BORDER = "D9E2F3"

OUT = os.path.join(os.path.dirname(__file__), "StudyFlow_AI_Submission.docx")


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def h1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = BRAND
    p.space_before = Pt(10)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = DARK
    return p


def body(doc, text, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = GREY
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = GREY
    return p


def kv_table(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for k, v in rows:
        r = t.add_row().cells
        shade(r[0], LIGHT_BG)
        set_cell_text(r[0], k, bold=True, color=DARK, size=10)
        set_cell_text(r[1], v, size=10, color=GREY)
        r[0].width = Inches(1.9)
        r[1].width = Inches(4.6)
    return t


def flow_row(doc, steps):
    """Render a horizontal flow: box -> box -> box using a single-row table."""
    cols = len(steps) * 2 - 1
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells = t.rows[0].cells
    for i, step in enumerate(steps):
        c = cells[i * 2]
        shade(c, LIGHT_BG)
        set_cell_text(c, step, bold=True, color=BRAND, size=9.5,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        if i < len(steps) - 1:
            arrow = cells[i * 2 + 1]
            set_cell_text(arrow, "→", bold=True, color=DARK, size=12,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    return t


def spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    return p


def build():
    doc = Document()
    # base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # ---------- Title ----------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("StudyFlow AI")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = BRAND
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Plan Smarter. Study Better.")
    rs.italic = True
    rs.font.size = Pt(13)
    rs.font.color.rgb = GREY
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = tag.add_run("Agentic AI Hackathon — Project Submission")
    rt.font.size = Pt(11)
    rt.font.color.rgb = DARK
    doc.add_paragraph()

    # ============================================================= 1
    h1(doc, "1. Project Overview")
    kv_table(doc, [
        ("Project Title", "StudyFlow AI — an adaptive, multi-agent academic planning assistant"),
        ("Tagline", "Plan Smarter. Study Better."),
        ("Team Members", "[Add your name(s) here]"),
        ("GitHub Repository", "https://github.com/sumukhveeramalla/studeyFlowAI"),
        ("Category", "Agentic AI — task planning, scheduling, validation, adaptive replanning"),
    ])
    spacer(doc)

    h2(doc, "Problem Statement")
    body(doc,
         "College students juggle assignments, exams, fixed classes, and personal "
         "commitments across disconnected tools (calendars, to-do apps, notes). "
         "Planning a realistic, conflict-free study schedule is tedious and brittle: "
         "when something unexpected comes up (a meeting, a missed session), students "
         "must manually re-plan everything while still respecting deadlines. Generic "
         "chatbots can discuss a plan but cannot reliably create items, compute "
         "availability, guarantee no scheduling conflicts, or safely change a "
         "calendar on the student's behalf.")

    h2(doc, "Target Users")
    bullet(doc, "College and university students managing multiple courses with assignments, exams, and deadlines.")
    bullet(doc, "Students who want to turn a free-form description of their workload into an organized, conflict-free study plan.")
    bullet(doc, "Learners who need the plan to adapt automatically when circumstances change — without losing control over their calendar.")

    h2(doc, "Proposed Solution")
    body(doc,
         "StudyFlow AI is a full-stack web application where the student talks to an "
         "AI assistant in natural language. The assistant extracts assignments, exams "
         "and tasks (asking follow-up questions for any missing deadline), creates "
         "them, auto-breaks assignments into subtasks, and generates a conflict-free "
         "study schedule using a deterministic scheduling engine. When the student "
         "reports an unexpected event, the system detects the conflict and proposes a "
         "revised schedule. Every schedule change is presented as a reviewable "
         "proposal and only applied after the student approves it. A live Agent "
         "Activity panel shows exactly which agents ran, which tools were called, and "
         "the validation result.")

    h2(doc, "Why Agentic AI Is Appropriate")
    bullet(doc, "The task needs interpretation of messy natural language, then a sequence of decisions (what to create, how to break it down, how to schedule, whether to re-plan) — a good fit for an agent that routes between specialized sub-tasks.", bold_lead="Conditional reasoning: ")
    bullet(doc, "Agents call real Python tools and database services (create items, compute availability, generate/validate schedules) rather than just producing text.", bold_lead="Tool use: ")
    bullet(doc, "Scheduling correctness (no overlaps, deadlines respected) is enforced by deterministic validation, not by the LLM — the agent recovers and retries when a proposal fails.", bold_lead="Verifiable actions: ")
    bullet(doc, "Changing a student's calendar is consequential, so the workflow pauses for explicit approval — a natural human-in-the-loop checkpoint.", bold_lead="Human-in-the-loop: ")

    doc.add_page_break()

    # ============================================================= 2
    h1(doc, "2. Agentic AI Design")

    h2(doc, "Architecture: Multi-Agent")
    body(doc,
         "StudyFlow AI uses a multi-agent architecture orchestrated with LangGraph. A "
         "Supervisor classifies the request and routes it to the right specialist; "
         "scheduling outputs pass through an independent, deterministic Validator that "
         "can send the workflow back for another attempt (a recovery loop). A separate, "
         "stateful Intake flow handles multi-turn item creation. The LLM is optional: "
         "with an OpenAI-compatible key the agents use real model reasoning; without "
         "one, the app runs a clearly-labeled deterministic demo mode.")

    h2(doc, "Overall Workflow")
    flow_row(doc, ["User message", "Supervisor\n(classify intent)", "Specialist\nagent + tools",
                   "Deterministic\nvalidation", "Proposal", "Human\napprove/reject", "Calendar\nupdated"])
    spacer(doc, 6)
    body(doc, "Intent routing from the Supervisor:")
    bullet(doc, "the message describes new assignments/exams/tasks → Intake agent (multi-turn slot-filling, then create + auto-decompose).", bold_lead="create: ")
    bullet(doc, "build a conflict-free study schedule → Scheduling agent → Validator → proposal.", bold_lead="schedule: ")
    bullet(doc, "an unexpected event changes things → Replanning agent → Validator → proposal.", bold_lead="replan: ")
    bullet(doc, "break an existing assignment into subtasks → Task-Planning agent.", bold_lead="decompose: ")
    bullet(doc, "answer questions about existing deadlines/priorities (read-only).", bold_lead="query: ")
    bullet(doc, "the request is too vague → ask the student for detail.", bold_lead="clarify: ")

    h2(doc, "Agent Responsibilities")
    rows = [
        ("Supervisor", "Interprets the request, classifies intent, extracts parameters, and routes to a specialist. Keyword router + LLM classifier, with unambiguous command phrases locked."),
        ("Intake agent", "Stateful, multi-turn. Parses a paragraph into assignment/exam/task drafts, asks follow-ups for missing required fields (title + deadline), summarizes for confirmation, then creates items on 'yes'."),
        ("Academic agent", "Reads courses/assignments/exams/tasks and answers queries; surfaces upcoming work ranked by deadline and priority."),
        ("Task-Planning agent", "Decomposes an assignment into ordered, validated subtasks with time estimates (LLM or deterministic template); totals reconcile to the assignment estimate to avoid double counting."),
        ("Scheduling agent", "Calls the deterministic engine to place study blocks around fixed commitments, honoring preferences and deadlines; reports infeasibility explicitly."),
        ("Replanning agent", "Injects the new event, detects which existing study blocks now conflict, preserves the rest, and reschedules only the affected work (minimal change)."),
        ("Validator", "Independent, deterministic checks (no LLM): no overlaps, no study block over a fixed event, start<end, within study hours, deadlines respected. On failure the graph retries (max 2) then finalizes with a feasibility note."),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    shade(hdr[0], LIGHT_BG); shade(hdr[1], LIGHT_BG)
    set_cell_text(hdr[0], "Agent", bold=True, color=DARK, size=10)
    set_cell_text(hdr[1], "Responsibility", bold=True, color=DARK, size=10)
    hdr[0].width = Inches(1.6); hdr[1].width = Inches(4.9)
    for name, desc in rows:
        c = t.add_row().cells
        set_cell_text(c[0], name, bold=True, color=BRAND, size=9.5)
        set_cell_text(c[1], desc, size=9.5, color=GREY)
        c[0].width = Inches(1.6); c[1].width = Inches(4.9)
    spacer(doc)

    h2(doc, "Tools & External Resources")
    bullet(doc, "get_academic_context, decompose_assignment, generate_schedule, replan_schedule, validate_schedule, list_events, get_preferences — thin Python functions over the database and scheduling/validation services.", bold_lead="Agent tools: ")
    bullet(doc, "A deterministic greedy scheduler computes free slots around fixed events and places study sessions (session length + breaks) before deadlines.", bold_lead="Scheduling engine: ")
    bullet(doc, "Optional OpenAI-compatible LLM via environment variables (tested with OpenRouter, e.g. openai/gpt-4o-mini) for intent classification, paragraph parsing, and decomposition. Falls back to deterministic logic if unavailable.", bold_lead="LLM: ")
    bullet(doc, "FastAPI + SQLAlchemy + SQLite (backend); React + TypeScript + Vite + Tailwind + Recharts (frontend); LangGraph for orchestration.", bold_lead="Frameworks: ")

    h2(doc, "Context / Information Required")
    bullet(doc, "Courses, assignments (title, due date, priority, estimated hours, status), exams (date, topics, prep hours), tasks (deadline, estimate, parent assignment), and calendar events (fixed vs movable study blocks).")
    bullet(doc, "Student preferences: study hours, session length, break length, available days, timezone — applied by the scheduler.")
    bullet(doc, "Shared LangGraph state per request: the message, detected intent, loaded context, generated proposal, validation result, and the ordered list of agent steps.")

    h2(doc, "Human-in-the-Loop Decisions")
    bullet(doc, "Item creation is confirmed in conversation: the assistant summarizes what it will create and only creates after the student says yes.")
    bullet(doc, "Schedule changes are never auto-applied. Generate/replan produce a pending proposal (original vs proposed, added/moved/removed blocks, explanations). The student clicks Approve or Reject; only Approve writes to the calendar.")

    doc.add_page_break()

    # ============================================================= 3
    h1(doc, "3. Working Prototype")
    body(doc,
         "The repository contains a fully functioning full-stack application. The main "
         "workflow — describe work in natural language → create items → generate a "
         "schedule → replan around an unexpected event → approve → persisted calendar "
         "— works end to end against real application state (not scripted responses).")

    h2(doc, "What Runs")
    bullet(doc, "Python 3.12, FastAPI, SQLAlchemy, SQLite, LangGraph; OpenAPI docs at /docs.", bold_lead="Backend: ")
    bullet(doc, "React + TypeScript + Vite + Tailwind + Recharts; 8 pages (Dashboard, Calendar, Assignments, Exams, Tasks, AI Assistant, Analytics, Settings).", bold_lead="Frontend: ")
    bullet(doc, "31 backend tests (pytest) covering CRUD/persistence, decomposition, overlap/deadline detection, scheduling feasibility, replan-after-conflict, approve/reject, API errors, multi-item paragraph parsing, and a full replanning smoke test; frontend Vitest tests + production build.", bold_lead="Tests: ")

    h2(doc, "How to Run")
    body(doc, "Backend:", size=10)
    code(doc, "cd backend\n"
              "uv venv --python 3.12 .venv\n"
              "uv pip install --python .venv/bin/python -r requirements.txt\n"
              "cp .env.example .env            # optional LLM key; works without one\n"
              ".venv/bin/python -m uvicorn app.main:app --reload --port 8000")
    body(doc, "Frontend (new terminal):", size=10)
    code(doc, "cd frontend\n"
              "npm install\n"
              "npm run dev                     # http://localhost:5173")
    body(doc, "Tests:", size=10)
    code(doc, "cd backend && .venv/bin/python -m pytest -q\n"
              "cd frontend && npm run test && npm run build")

    h2(doc, "LLM Configuration (optional)")
    body(doc,
         "Set OPENAI_API_KEY in backend/.env (gitignored). For OpenRouter, set "
         "OPENAI_BASE_URL=https://openrouter.ai/api/v1 and LLM_MODEL=openai/gpt-4o-mini. "
         "GET /api/health reports {\"llm_enabled\": true} when a key is active. With no "
         "key, the app runs a clearly-labeled deterministic demo mode.")

    doc.add_page_break()

    # ============================================================= 4
    h1(doc, "4. Demo — Problem → Interaction → Agentic Workflow → Result")

    h2(doc, "Scenario")
    body(doc,
         "A student types a single paragraph describing everything due, then asks for "
         "a plan; later an unexpected meeting appears and the plan adapts.")

    h2(doc, "Step-by-step")
    demo = [
        ("Problem", "The student has several deadlines (e.g. an assignment due tonight, a presentation to prepare for tomorrow, and a paper revision due next week) and no organized plan."),
        ("User interaction", "In the AI Assistant, the student writes it all in one message: \"I need to submit X tonight, prepare the CS670 presentation tomorrow at 4:30pm (4 hrs), and write a response letter and submit before Nov 4.\""),
        ("Agentic workflow — create", "Supervisor routes to the Intake agent. It parses the paragraph into distinct items with deadlines, asks for any missing deadline, summarizes, and on 'yes' creates them — auto-decomposing each assignment into subtasks."),
        ("Agentic workflow — schedule", "The student says \"create a study plan for this week.\" The Scheduling agent places study blocks around fixed events; the Validator confirms no conflicts; a proposal is returned (e.g. several new study blocks)."),
        ("Agentic workflow — replan", "The student says \"I have an unexpected meeting tomorrow 9–11am, rearrange my schedule.\" The Replanning agent detects conflicting blocks, reschedules the affected work, and the Validator checks the result."),
        ("Human-in-the-loop", "Each proposal shows original vs proposed blocks with explanations and Approve / Reject buttons. Nothing changes until the student approves."),
        ("Result", "On Approve, the calendar, dashboard statistics, and task progress update immediately and persist across refresh. The Agent Activity panel shows the exact agents, tools, and validation outcome for the run."),
    ]
    for label, text in demo:
        p = doc.add_paragraph()
        rl = p.add_run(label + ": ")
        rl.bold = True
        rl.font.color.rgb = BRAND
        rl.font.size = Pt(10.5)
        rt = p.add_run(text)
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = GREY

    body(doc,
         "Core value shown: the assistant turns messy natural language into real, "
         "persisted academic items and a conflict-free, deadline-aware schedule, and "
         "adapts safely under human approval — genuine agentic planning, not a chatbot.")

    doc.add_page_break()

    # ============================================================= 5
    h1(doc, "5. Evaluation")
    body(doc,
         "The following representative cases are backed by the automated test suite "
         "(31 passing backend tests) and by live runs against a real LLM via OpenRouter.")

    cases = [
        ("TC-1 Multi-item paragraph intake",
         "One paragraph with an assignment, an exam and a task is split into 3 distinct items with the right types and deadlines.",
         "Parsed into assignment + exam + task; exam topics captured; each item got a date. (test_parse_multi_item_paragraph, test_multi_item_paragraph_end_to_end)",
         "Pass"),
        ("TC-2 Missing-deadline follow-up",
         "If an item has no date, the assistant asks for it, then creates after confirmation.",
         "Assistant asked for the due date, accepted a natural-language reply, confirmed, and created the item with linked subtasks. (test_intake_*)",
         "Pass"),
        ("TC-3 Conflict-free scheduling",
         "Generated study blocks never overlap fixed events or each other and respect deadlines.",
         "Validator confirmed zero overlaps / fixed conflicts; blocks placed before deadlines. (test_scheduling)",
         "Pass"),
        ("TC-4 Infeasibility reporting",
         "When there is not enough time before a deadline, the system reports a shortfall instead of inventing slots.",
         "Returned an explicit infeasible item (needed vs placed minutes). (test_schedule_infeasible_reports_shortfall)",
         "Pass"),
        ("TC-5 Dynamic replanning",
         "An unexpected event triggers detection of conflicting blocks, rescheduling, validation, and a proposal.",
         "Supervisor→Replanner→Validator→await-approval; conflicting block removed, work rescheduled, no overlap with the new event. (test_full_replanning_demo_smoke)",
         "Pass"),
        ("TC-6 Human approval / rejection",
         "Approve applies changes and persists them; Reject changes nothing; a decided proposal can't be decided again.",
         "Approve persisted the new meeting and blocks; Reject left event counts unchanged; second decision returned HTTP 409. (test_reject_proposal_has_no_effect)",
         "Pass"),
        ("TC-7 Intent routing (regression fix)",
         "A message that describes new work AND says 'plan my schedule' should create first; 'create a study plan' alone should schedule.",
         "Paragraph routed to create; 'create a study plan' routed to schedule and produced real study blocks. Verified live.",
         "Pass"),
    ]
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, label in enumerate(["Test case", "Expected", "Actual", "Result"]):
        shade(hdr[i], LIGHT_BG)
        set_cell_text(hdr[i], label, bold=True, color=DARK, size=9.5)
    for name, exp, act, res in cases:
        c = t.add_row().cells
        set_cell_text(c[0], name, bold=True, color=BRAND, size=9)
        set_cell_text(c[1], exp, size=9, color=GREY)
        set_cell_text(c[2], act, size=9, color=GREY)
        set_cell_text(c[3], res, bold=True, color=RGBColor(0x05, 0x96, 0x69), size=9)
        c[0].width = Inches(1.4); c[1].width = Inches(1.9)
        c[2].width = Inches(2.6); c[3].width = Inches(0.6)
    spacer(doc)

    h2(doc, "Successful Behaviors")
    bullet(doc, "Natural-language paragraphs become real, persisted items with deadlines and auto-generated subtasks.")
    bullet(doc, "Deterministic validation guarantees conflict-free schedules; the agent recovers and retries on failure.")
    bullet(doc, "Human-in-the-loop approval: no calendar change happens without explicit consent; changes persist after refresh.")
    bullet(doc, "Observable agent activity (agent, tool, validation) reflects real execution — nothing fabricated.")

    h2(doc, "Failure Cases / Limitations")
    bullet(doc, "In deterministic demo mode (no API key), paragraph parsing is heuristic: titles can keep small artifacts and compound items (\"write a letter and change the paper\") may merge or default their estimate.")
    bullet(doc, "The LLM occasionally types a 'prepare for presentation' item as a task rather than an assignment, and may not capture a stated prep-hours estimate (a default is used).")
    bullet(doc, "The scheduler is a documented greedy MVP (earliest-deadline, then priority): conflict-free but not globally optimal.")
    bullet(doc, "Single demo-student profile, no authentication. Timezone is single-user (stored UTC, rendered local).")
    bullet(doc, "Canvas / Google Calendar / notifications are intentionally not implemented (listed as stretch goals).")

    h2(doc, "Possible Improvements")
    bullet(doc, "Richer LLM extraction prompt to capture exam time-of-day and prep hours and to split compound items reliably.")
    bullet(doc, "Optimization-based scheduler (e.g. weighted constraint solver) and spaced-repetition for exam prep.")
    bullet(doc, "Multi-user accounts with authentication and real Canvas / Google Calendar sync.")
    bullet(doc, "Learned personalization of productive hours; proactive reminders and notifications.")

    doc.add_page_break()

    # ============================================================= 6
    h1(doc, "6. Project Documentation")

    h2(doc, "What Was Built")
    body(doc,
         "A working full-stack, multi-agent academic assistant: a React/TypeScript "
         "frontend (8 pages) over a FastAPI/SQLAlchemy/SQLite backend, with a "
         "LangGraph agent workflow, a deterministic scheduling engine, an independent "
         "validation layer, and a human-approval proposal system.")

    h2(doc, "How the System Works")
    bullet(doc, "The frontend calls typed REST endpoints; TanStack Query keeps Dashboard/Calendar/Assignments/Tasks/Exams in sync.", bold_lead="Frontend → API: ")
    bullet(doc, "Routers → services → ORM. Services include scheduling, validation, proposals, decomposition, intake (multi-turn), and dashboard stats.", bold_lead="Backend layers: ")
    bullet(doc, "A Supervisor routes intent; specialists call tools; a Validator enforces correctness with a bounded recovery loop; a Finalize node creates a pending proposal. Every step is logged for the Agent Activity panel.", bold_lead="Agents: ")
    bullet(doc, "Generate/replan produce proposals; Approve applies them atomically and persists; Reject discards them.", bold_lead="Approval: ")

    h2(doc, "How to Run / Access")
    body(doc,
         "See Section 3. In short: start the backend (uvicorn on :8000) and the "
         "frontend (Vite on :5173). The app runs with no API key (demo mode) or with "
         "an OpenAI-compatible key in backend/.env for full LLM reasoning. Interactive "
         "API docs are available at /docs.")

    h2(doc, "Major Tools & Technologies")
    kv_table(doc, [
        ("Backend", "Python 3.12, FastAPI, Pydantic, SQLAlchemy, SQLite, Uvicorn"),
        ("Agents", "LangGraph, LangChain-OpenAI (OpenAI-compatible; tested on OpenRouter)"),
        ("Frontend", "React, TypeScript, Vite, Tailwind CSS, Recharts, lucide-react, TanStack Query"),
        ("Testing", "pytest (31 tests), Vitest + Testing Library"),
        ("Repository", "https://github.com/sumukhveeramalla/studeyFlowAI"),
    ])
    spacer(doc)

    h2(doc, "Known Limitations")
    bullet(doc, "Demo-mode parsing is heuristic; LLM mode is markedly better for messy input.")
    bullet(doc, "Greedy (not optimal) scheduler; single-user, no authentication; single timezone.")
    bullet(doc, "No external integrations (Canvas, Google Calendar, notifications) in this version.")
    bullet(doc, "A stale local server can hold port 8000 between runs; free it before restarting.")

    doc.save(OUT)
    print("Wrote", OUT)


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = DARK
    # light shading via a bordered single-cell would be heavier; keep monospace.
    return p


if __name__ == "__main__":
    build()
